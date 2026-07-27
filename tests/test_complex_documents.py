import io
import unittest

from docx import Document
from openpyxl import Workbook

from src.io import LoaderOptions, load_documents_from_bytes


def make_minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 24 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]

    pdf = io.BytesIO()
    pdf.write(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(pdf.tell())
        pdf.write(f"{index} 0 obj\n".encode("ascii"))
        pdf.write(obj)
        pdf.write(b"\nendobj\n")

    xref_offset = pdf.tell()
    pdf.write(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.write(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return pdf.getvalue()


class ComplexDocumentLoaderTests(unittest.TestCase):
    def test_pdf_loader_extracts_page_text(self):
        pdf_bytes = make_minimal_pdf("Section One\nPDF alpha document")

        texts, metadatas = load_documents_from_bytes("sample.pdf", pdf_bytes)

        self.assertEqual(len(texts), 1)
        self.assertIn("PDF alpha document", texts[0])
        self.assertEqual(metadatas[0]["page_number"], 1)
        self.assertEqual(metadatas[0]["document_part"], "page")
        self.assertEqual(metadatas[0]["section_title"], "Section One")
        self.assertFalse(metadatas[0]["pdf_needs_ocr"])

    def test_pdf_loader_reports_scanned_pdf_when_no_text(self):
        pdf_bytes = make_minimal_pdf("")

        with self.assertRaises(ValueError):
            load_documents_from_bytes("scan.pdf", pdf_bytes)

    def test_pdf_loader_uses_static_ocr_when_enabled(self):
        pdf_bytes = make_minimal_pdf("")

        texts, metadatas = load_documents_from_bytes(
            "scan.pdf",
            pdf_bytes,
            options=LoaderOptions(ocr_enabled=True, ocr_static_pages={1: "OCR alpha text"}),
        )

        self.assertEqual(texts, ["OCR alpha text"])
        self.assertTrue(metadatas[0]["ocr_applied"])
        self.assertEqual(metadatas[0]["ocr_backend"], "tesseract")
        self.assertEqual(metadatas[0]["ocr_language"], "eng")

    def test_image_loader_uses_static_ocr_when_enabled(self):
        texts, metadatas = load_documents_from_bytes(
            "sample.png",
            b"not a real image",
            options=LoaderOptions(ocr_enabled=True, ocr_static_pages={1: "image OCR text"}),
        )

        self.assertEqual(texts, ["image OCR text"])
        self.assertEqual(metadatas[0]["document_part"], "image")
        self.assertTrue(metadatas[0]["ocr_applied"])

    def test_docx_loader_extracts_heading_sections_and_tables(self):
        document = Document()
        document.add_heading("Chapter One", level=1)
        document.add_paragraph("DOCX alpha paragraph")
        document.add_heading("Details", level=2)
        document.add_paragraph("Nested detail paragraph")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header A"
        table.cell(0, 1).text = "Header B"
        table.cell(1, 0).text = "Cell A"
        table.cell(1, 1).text = "Cell B"
        buffer = io.BytesIO()
        document.save(buffer)

        texts, metadatas = load_documents_from_bytes("sample.docx", buffer.getvalue())

        self.assertEqual(len(texts), 2)
        self.assertIn("DOCX alpha paragraph", texts[0])
        self.assertIn("Cell A", texts[1])
        self.assertEqual(metadatas[0]["file_type"], ".docx")
        self.assertEqual(metadatas[0]["section_title"], "Chapter One")
        self.assertEqual(metadatas[1]["heading_path"], "Chapter One > Details")
        self.assertEqual(metadatas[1]["table_count"], 1)

    def test_xlsx_loader_extracts_sheet_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Name", "Value"])
        sheet.append(["alpha", 42])
        buffer = io.BytesIO()
        workbook.save(buffer)

        texts, metadatas = load_documents_from_bytes("sample.xlsx", buffer.getvalue())

        self.assertEqual(len(texts), 1)
        self.assertIn("Name: alpha", texts[0])
        self.assertIn("Value: 42", texts[0])
        self.assertEqual(metadatas[0]["sheet_name"], "Data")
        self.assertEqual(metadatas[0]["document_part"], "excel_row")
        self.assertEqual(metadatas[0]["row_number"], 2)

    def test_xlsx_loader_combines_merged_multirow_headers(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Merged"
        sheet.merge_cells("A1:B1")
        sheet["A1"] = "Region"
        sheet.append(["East", "West"])
        sheet.append([10, 20])
        buffer = io.BytesIO()
        workbook.save(buffer)

        texts, metadatas = load_documents_from_bytes(
            "merged.xlsx",
            buffer.getvalue(),
            options=LoaderOptions(excel_header_rows=2),
        )

        self.assertEqual(len(texts), 1)
        self.assertIn("Region / East: 10", texts[0])
        self.assertIn("Region / West: 20", texts[0])
        self.assertEqual(metadatas[0]["header_row_start"], 1)
        self.assertEqual(metadatas[0]["header_row_end"], 2)
        self.assertIn("A1:B1", metadatas[0]["merged_ranges"])

    def test_xlsx_loader_respects_record_limit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name"])
        sheet.append(["row one"])
        sheet.append(["row two"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        texts, metadatas = load_documents_from_bytes(
            "limited.xlsx",
            buffer.getvalue(),
            options=LoaderOptions(max_records=1),
        )

        self.assertIn("row one", texts[0])
        self.assertEqual(len(texts), 1)
        self.assertNotIn("row two", texts[0])
        self.assertEqual(metadatas[0]["row_count"], 1)


if __name__ == "__main__":
    unittest.main()
