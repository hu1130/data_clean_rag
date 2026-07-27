"""文件加载器，统一把不同格式转换为文本和元数据。"""

import csv
import importlib.metadata
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from src.ingestion import compute_content_hash, file_key_for


DEFAULT_TEXT_FIELDS = (
    "text",
    "content",
    "body",
    "page_content",
    "description",
    "summary",
    "title",
)

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".jsonl",
    ".csv",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
}

TEXT_EXTENSIONS = {".txt", ".md", ".json", ".jsonl", ".csv", ".html", ".htm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
PROJECT_ROOT = Path(__file__).resolve().parents[2]


@dataclass
class LoaderOptions:
    """文件加载配置。"""

    text_fields: Optional[Iterable[str]] = None
    max_file_size_mb: int = 100
    max_records: Optional[int] = None
    encoding: str = "utf-8"
    keep_table_markdown: bool = True
    excel_sheet_limit: Optional[int] = None
    excel_header_row: Optional[int] = None
    excel_header_rows: int = 1
    excel_row_documents: bool = True
    docx_split_by_heading: bool = True
    pdf_infer_sections: bool = True
    ocr_enabled: bool = False
    ocr_backend: str = "tesseract"
    ocr_language: str = "eng"
    ocr_dpi: int = 200
    ocr_max_pages: Optional[int] = None
    ocr_static_pages: Optional[Dict[int, str]] = None

    @property
    def max_file_size_bytes(self) -> int:
        return self.max_file_size_mb * 1024 * 1024


def load_documents_from_path(
    path: str,
    text_fields: Optional[Iterable[str]] = None,
    options: Optional[LoaderOptions] = None
) -> Tuple[List[str], List[Dict[str, Any]]]:
    """从文件或目录加载文本。"""
    loader_options = _resolve_options(options, text_fields)
    source_path = Path(path)
    if source_path.is_dir():
        texts: List[str] = []
        metadatas: List[Dict[str, Any]] = []
        for file_path in sorted(source_path.rglob("*")):
            if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
                file_texts, file_metas = load_documents_from_path(str(file_path), options=loader_options)
                texts.extend(file_texts)
                metadatas.extend(file_metas)
        return texts, metadatas

    _validate_file_size(source_path.stat().st_size, loader_options)
    raw = source_path.read_bytes()
    return load_documents_from_bytes(
        filename=source_path.name,
        content=raw,
        options=loader_options,
        base_metadata={
            "source_file": str(source_path),
            "file_type": source_path.suffix.lower(),
        },
    )


def load_documents_from_bytes(
    filename: str,
    content: bytes,
    text_fields: Optional[Iterable[str]] = None,
    base_metadata: Optional[Dict[str, Any]] = None,
    options: Optional[LoaderOptions] = None,
) -> Tuple[List[str], List[Dict[str, Any]]]:
    """从上传文件内容加载文本。"""
    loader_options = _resolve_options(options, text_fields)
    _validate_file_size(len(content), loader_options)

    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {extension}")

    metadata = dict(base_metadata or {})
    metadata.setdefault("filename", filename)
    metadata.setdefault("file_type", extension)
    content_hash = compute_content_hash(content)
    metadata.setdefault("file_hash", content_hash)
    metadata.setdefault("content_hash", content_hash)
    metadata.setdefault("file_size_bytes", len(content))
    metadata.setdefault("file_key", file_key_for(filename, metadata.get("source_file")))

    if extension in {".txt", ".md"}:
        return _single_document(_decode_text(content, loader_options), metadata)
    if extension == ".json":
        return _load_json(_decode_text(content, loader_options), metadata, loader_options)
    if extension == ".jsonl":
        return _load_jsonl(_decode_text(content, loader_options), metadata, loader_options)
    if extension == ".csv":
        return _load_csv(_decode_text(content, loader_options), metadata, loader_options)
    if extension in {".html", ".htm"}:
        return _load_html(_decode_text(content, loader_options), metadata)
    if extension == ".pdf":
        return _load_pdf(content, metadata, loader_options)
    if extension == ".docx":
        return _load_docx(content, metadata, loader_options)
    if extension in {".xlsx", ".xlsm"}:
        return _load_xlsx(content, metadata, loader_options)
    if extension in IMAGE_EXTENSIONS:
        return _load_image_ocr(content, metadata, loader_options)

    return _single_document(_decode_text(content, loader_options), metadata)


def _resolve_options(
    options: Optional[LoaderOptions],
    text_fields: Optional[Iterable[str]]
) -> LoaderOptions:
    if options is None:
        return LoaderOptions(text_fields=text_fields)

    if text_fields is not None:
        options = LoaderOptions(
            text_fields=text_fields,
            max_file_size_mb=options.max_file_size_mb,
            max_records=options.max_records,
            encoding=options.encoding,
            keep_table_markdown=options.keep_table_markdown,
            excel_sheet_limit=options.excel_sheet_limit,
            excel_header_row=options.excel_header_row,
            excel_header_rows=options.excel_header_rows,
            excel_row_documents=options.excel_row_documents,
            docx_split_by_heading=options.docx_split_by_heading,
            pdf_infer_sections=options.pdf_infer_sections,
            ocr_enabled=options.ocr_enabled,
            ocr_backend=options.ocr_backend,
            ocr_language=options.ocr_language,
            ocr_dpi=options.ocr_dpi,
            ocr_max_pages=options.ocr_max_pages,
            ocr_static_pages=options.ocr_static_pages,
        )
    return options


def _validate_file_size(size_bytes: int, options: LoaderOptions):
    if size_bytes > options.max_file_size_bytes:
        size_mb = size_bytes / 1024 / 1024
        raise ValueError(f"文件过大: {size_mb:.1f}MB，限制为 {options.max_file_size_mb}MB")


def _decode_text(content: bytes, options: LoaderOptions) -> str:
    try:
        return content.decode(options.encoding)
    except UnicodeDecodeError:
        return content.decode("utf-8-sig", errors="replace")


def _single_document(text: str, metadata: Dict[str, Any]) -> Tuple[List[str], List[Dict[str, Any]]]:
    clean_text = text.strip()
    if not clean_text:
        return [], []
    return [clean_text], [metadata]


def _load_json(
    text: str,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    data = json.loads(text)
    records = data if isinstance(data, list) else [data]
    return _records_to_documents(_limit_records(records, options), metadata, options)


def _load_jsonl(
    text: str,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    records = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        if options.max_records is not None and len(records) >= options.max_records:
            break

        line = line.strip()
        if not line:
            continue
        record = json.loads(line)
        if isinstance(record, dict):
            record = {**record, "_line_number": line_number}
        records.append(record)
    return _records_to_documents(records, metadata, options)


def _load_csv(
    text: str,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    reader = csv.DictReader(text.splitlines())
    records = []
    for row_number, row in enumerate(reader, start=1):
        if options.max_records is not None and len(records) >= options.max_records:
            break
        records.append({**row, "_row_number": row_number})
    return _records_to_documents(records, metadata, options)


def _load_html(text: str, metadata: Dict[str, Any]) -> Tuple[List[str], List[Dict[str, Any]]]:
    clean_text = _html_to_text(text)
    return _single_document(clean_text, metadata)


def _load_pdf(
    content: bytes,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError("处理 PDF 需要安装 pypdf: pip install pypdf")

    reader = PdfReader(io.BytesIO(content))
    extracted_pages: Dict[int, str] = {}
    current_section: Optional[str] = None
    empty_page_numbers: List[int] = []
    for page_index, page in enumerate(reader.pages, start=1):
        if options.max_records is not None and len(extracted_pages) >= options.max_records:
            break

        page_text = (page.extract_text() or "").strip()
        if not page_text:
            empty_page_numbers.append(page_index)
            continue

        extracted_pages[page_index] = page_text

    ocr_pages: Dict[int, str] = {}
    if empty_page_numbers and options.ocr_enabled:
        ocr_pages = _ocr_pdf_pages(content, empty_page_numbers, options)

    all_page_texts = {**extracted_pages, **ocr_pages}
    if not all_page_texts and empty_page_numbers:
        if options.ocr_enabled:
            raise ValueError("PDF OCR did not produce extractable text.")
        raise ValueError("PDF has no extractable text; it is likely scanned and requires OCR.")

    texts: List[str] = []
    metadatas: List[Dict[str, Any]] = []
    empty_pages = len(empty_page_numbers)
    for page_index in sorted(all_page_texts):
        if options.max_records is not None and len(texts) >= options.max_records:
            break

        page_text = all_page_texts[page_index].strip()
        if not page_text:
            continue

        if options.pdf_infer_sections:
            inferred_section = _infer_section_title(page_text)
            if inferred_section:
                current_section = inferred_section

        page_meta = dict(metadata)
        page_meta["page_number"] = page_index
        page_meta["total_pages"] = len(reader.pages)
        page_meta["document_part"] = "page"
        page_meta["ocr_applied"] = page_index in ocr_pages
        if page_index in ocr_pages:
            page_meta["ocr_backend"] = options.ocr_backend
            page_meta["ocr_language"] = options.ocr_language
        if current_section:
            page_meta["section_title"] = current_section
            page_meta["section_path"] = current_section
        texts.append(page_text)
        metadatas.append(page_meta)

    for page_meta in metadatas:
        page_meta["pdf_text_pages"] = len(texts)
        page_meta["pdf_empty_pages"] = empty_pages
        page_meta["pdf_needs_ocr"] = empty_pages > 0

    return texts, metadatas


def _load_docx(
    content: bytes,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise ImportError("处理 DOCX 需要安装 python-docx: pip install python-docx")

    document = Document(io.BytesIO(content))

    if not options.docx_split_by_heading:
        parts = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        if options.keep_table_markdown:
            for table_index, table in enumerate(document.tables, start=1):
                table_text = _docx_table_to_markdown(table, table_index)
                if table_text:
                    parts.append(table_text)

        return _single_document("\n\n".join(parts), metadata)

    texts: List[str] = []
    metadatas: List[Dict[str, Any]] = []
    current_parts: List[str] = []
    current_title: Optional[str] = None
    current_level: Optional[int] = None
    current_table_count = 0
    heading_stack: List[Optional[str]] = []
    table_index = 0

    def flush_section():
        nonlocal current_parts, current_title, current_level, current_table_count
        if not current_parts:
            return
        if options.max_records is not None and len(texts) >= options.max_records:
            current_parts = []
            return

        section_text = "\n\n".join(current_parts).strip()
        if not section_text:
            current_parts = []
            return

        section_meta = dict(metadata)
        section_meta["document_part"] = "section" if current_title else "document"
        section_meta["section_index"] = len(texts) + 1
        section_meta["table_count"] = current_table_count
        if current_title:
            section_meta["section_title"] = current_title
            section_meta["heading_level"] = current_level
            section_meta["heading_path"] = " > ".join([item for item in heading_stack if item])

        texts.append(section_text)
        metadatas.append(section_meta)
        current_parts = []
        current_table_count = 0

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue

            heading_level = _docx_heading_level(paragraph)
            if heading_level is not None:
                if len(current_parts) > 1 or (current_parts and not current_title):
                    flush_section()
                else:
                    current_parts = []
                    current_table_count = 0

                while len(heading_stack) < heading_level:
                    heading_stack.append(None)
                heading_stack = heading_stack[:heading_level]
                heading_stack[heading_level - 1] = text
                current_title = text
                current_level = heading_level
                current_parts = [text]
                continue

            current_parts.append(text)

        elif isinstance(child, CT_Tbl) and options.keep_table_markdown:
            table_index += 1
            table = Table(child, document)
            table_text = _docx_table_to_markdown(table, table_index)
            if table_text:
                current_parts.append(table_text)
                current_table_count += 1

    flush_section()
    return texts, metadatas


def _load_xlsx(
    content: bytes,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("处理 Excel 需要安装 openpyxl: pip install openpyxl")

    workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    texts: List[str] = []
    metadatas: List[Dict[str, Any]] = []

    sheet_names = workbook.sheetnames
    if options.excel_sheet_limit is not None:
        sheet_names = sheet_names[:options.excel_sheet_limit]

    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        rows = _excel_rows_with_merged_values(sheet)
        if not rows:
            continue

        header_start = _excel_header_start_row(rows, options)
        if header_start is None:
            continue

        header_row_count = max(1, options.excel_header_rows)
        header_end = min(header_start + header_row_count - 1, len(rows))
        headers = _combine_excel_headers(rows[header_start - 1:header_end])
        merged_ranges = _excel_merged_ranges(sheet)
        emitted_for_sheet = 0
        sheet_texts_start = len(texts)

        if options.excel_row_documents:
            for row_number in range(header_end + 1, len(rows) + 1):
                if options.max_records is not None and emitted_for_sheet >= options.max_records:
                    break

                row_values = rows[row_number - 1]
                if not any(row_values):
                    continue

                row_text = _excel_row_to_text(sheet_name, row_number, headers, row_values)
                if not row_text:
                    continue

                row_meta = dict(metadata)
                row_meta["document_part"] = "excel_row"
                row_meta["sheet_name"] = sheet_name
                row_meta["row_number"] = row_number
                row_meta["header_row_start"] = header_start
                row_meta["header_row_end"] = header_end
                row_meta["headers"] = " | ".join(headers)
                row_meta["merged_ranges"] = ";".join(merged_ranges)
                row_meta["record_index"] = len(texts)
                texts.append(row_text)
                metadatas.append(row_meta)
                emitted_for_sheet += 1
        else:
            sheet_rows = []
            for row_number, row_values in enumerate(rows, start=1):
                if options.max_records is not None and len(sheet_rows) >= options.max_records:
                    break
                values = [value for value in row_values if value]
                if values:
                    sheet_rows.append(" | ".join(values))

            if sheet_rows:
                sheet_meta = dict(metadata)
                sheet_meta["document_part"] = "excel_sheet"
                sheet_meta["sheet_name"] = sheet_name
                sheet_meta["row_count"] = len(sheet_rows)
                sheet_meta["header_row_start"] = header_start
                sheet_meta["header_row_end"] = header_end
                sheet_meta["headers"] = " | ".join(headers)
                sheet_meta["merged_ranges"] = ";".join(merged_ranges)
                texts.append("\n".join(sheet_rows))
                metadatas.append(sheet_meta)
                emitted_for_sheet = len(sheet_rows)

        for meta in metadatas[sheet_texts_start:]:
            meta["row_count"] = emitted_for_sheet

    workbook.close()
    return texts, metadatas


def _records_to_documents(
    records: List[Any],
    base_metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    texts: List[str] = []
    metadatas: List[Dict[str, Any]] = []

    for index, record in enumerate(records):
        text, field_names = _extract_text(record, options.text_fields)
        text = text.strip()
        if not text:
            continue

        metadata = dict(base_metadata)
        metadata["record_index"] = index
        if field_names:
            metadata["text_fields"] = ",".join(field_names)

        if isinstance(record, dict):
            for key in ("id", "_line_number", "_row_number"):
                if key in record:
                    metadata[key.lstrip("_")] = record[key]

        texts.append(text)
        metadatas.append(metadata)

    return texts, metadatas


def _limit_records(records: List[Any], options: LoaderOptions) -> List[Any]:
    if options.max_records is None:
        return records
    return records[:options.max_records]


def _extract_text(record: Any, text_fields: Optional[Iterable[str]]) -> Tuple[str, List[str]]:
    fields = tuple(text_fields or DEFAULT_TEXT_FIELDS)

    if isinstance(record, str):
        return record, []
    if isinstance(record, (int, float, bool)):
        return str(record), []
    if isinstance(record, list):
        parts = []
        used_fields: List[str] = []
        for item in record:
            item_text, item_fields = _extract_text(item, fields)
            if item_text:
                parts.append(item_text)
                used_fields.extend(item_fields)
        return "\n".join(parts), used_fields
    if not isinstance(record, dict):
        return "", []

    parts = []
    used_fields = []
    for field in fields:
        value = _get_nested_value(record, field)
        if value is None:
            continue
        value_text = str(value).strip()
        if value_text:
            parts.append(value_text)
            used_fields.append(field)

    if parts:
        return "\n".join(parts), used_fields

    fallback_parts = [
        str(value).strip()
        for key, value in record.items()
        if not key.startswith("_") and isinstance(value, (str, int, float, bool)) and str(value).strip()
    ]
    return "\n".join(fallback_parts), []


def _get_nested_value(record: Dict[str, Any], field: str) -> Any:
    if field in record:
        return record[field]

    current: Any = record
    for part in field.split("."):
        if not isinstance(current, dict) or part not in current:
            return None
        current = current[part]
    return current


def _html_to_text(html: str) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return _basic_html_to_text(html)

    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    return soup.get_text(separator="\n", strip=True)


def _basic_html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style|noscript).*?>.*?</\1>", " ", html)
    html = re.sub(r"(?s)<[^>]+>", " ", html)
    html = re.sub(r"\s+", " ", html)
    return html.strip()


def _load_image_ocr(
    content: bytes,
    metadata: Dict[str, Any],
    options: LoaderOptions
) -> Tuple[List[str], List[Dict[str, Any]]]:
    if not options.ocr_enabled:
        raise ValueError("Image files require OCR. Set ocr_enabled=true.")

    page_texts = _ocr_image(content, options)
    text = page_texts.get(1, "").strip()
    if not text:
        raise ValueError("Image OCR did not produce extractable text.")

    image_meta = dict(metadata)
    image_meta["document_part"] = "image"
    image_meta["ocr_applied"] = True
    image_meta["ocr_backend"] = options.ocr_backend
    image_meta["ocr_language"] = options.ocr_language
    return [text], [image_meta]


def _ocr_tool_bins() -> List[Path]:
    bins: List[Path] = []
    for env_name in ("OCR_TOOL_BIN", "POPPLER_PATH"):
        env_value = os.environ.get(env_name)
        if env_value:
            bins.append(Path(env_value))

    bins.extend([
        PROJECT_ROOT / "e" / "Library" / "bin",
        PROJECT_ROOT / ".m" / "e" / "Library" / "bin",
        PROJECT_ROOT / ".tools" / "ocr-env" / "Library" / "bin",
        PROJECT_ROOT / ".tools" / "Tesseract-OCR",
    ])
    return bins


def _resolve_tesseract_cmd() -> Optional[str]:
    tesseract_cmd = os.environ.get("TESSERACT_CMD")
    if not tesseract_cmd:
        for bin_dir in _ocr_tool_bins():
            candidate = bin_dir / ("tesseract.exe" if os.name == "nt" else "tesseract")
            if candidate.exists():
                tesseract_cmd = str(candidate)
                break

    if not tesseract_cmd:
        tesseract_cmd = shutil.which("tesseract")
    return tesseract_cmd


def _configure_tesseract(pytesseract: Any):
    tesseract_cmd = _resolve_tesseract_cmd()
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        tessdata_path = _resolve_tessdata_path(Path(tesseract_cmd))
        if tessdata_path:
            os.environ.setdefault("TESSDATA_PREFIX", tessdata_path)


def _resolve_tessdata_path(tesseract_cmd: Optional[Path] = None) -> Optional[str]:
    env_value = os.environ.get("TESSDATA_PREFIX")
    if env_value:
        return env_value

    candidates = [
        PROJECT_ROOT / "e" / "share" / "tessdata",
        PROJECT_ROOT / ".m" / "e" / "share" / "tessdata",
        PROJECT_ROOT / ".tools" / "ocr-env" / "share" / "tessdata",
    ]
    if tesseract_cmd:
        candidates.extend([
            tesseract_cmd.parent / "tessdata",
            tesseract_cmd.parent.parent.parent / "share" / "tessdata",
        ])

    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return None


def _resolve_poppler_path() -> Optional[str]:
    env_value = os.environ.get("POPPLER_PATH")
    if env_value:
        return env_value

    executable = "pdftoppm.exe" if os.name == "nt" else "pdftoppm"
    for bin_dir in _ocr_tool_bins():
        if (bin_dir / executable).exists():
            return str(bin_dir)

    system_executable = shutil.which(executable)
    if system_executable:
        return str(Path(system_executable).parent)

    return None


def _dependency_status(distribution: str, module: str) -> Dict[str, Any]:
    available = importlib.util.find_spec(module) is not None
    version = None
    if available:
        try:
            version = importlib.metadata.version(distribution)
        except importlib.metadata.PackageNotFoundError:
            version = None
    return {
        "available": available,
        "module": module,
        "version": version,
    }


def _executable_version(executable: Optional[str], *args: str) -> Dict[str, Any]:
    if not executable:
        return {
            "available": False,
            "path": None,
            "version": None,
            "error": "Executable not found",
        }

    try:
        result = subprocess.run(
            [executable, *args],
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )
    except Exception as exc:
        return {
            "available": False,
            "path": executable,
            "version": None,
            "error": str(exc),
        }

    output = (result.stdout or result.stderr or "").strip()
    return {
        "available": result.returncode == 0,
        "path": executable,
        "version": output.splitlines()[0] if output else None,
        "error": None if result.returncode == 0 else output,
    }


def _available_tessdata_languages(tessdata_path: Optional[str]) -> List[str]:
    if not tessdata_path:
        return []
    path = Path(tessdata_path)
    if not path.exists():
        return []
    return sorted(file_path.stem for file_path in path.glob("*.traineddata"))


def get_loader_diagnostics() -> Dict[str, Any]:
    """Return file loading, optional dependency, and OCR readiness diagnostics."""
    dependencies = {
        "pypdf": _dependency_status("pypdf", "pypdf"),
        "python-docx": _dependency_status("python-docx", "docx"),
        "openpyxl": _dependency_status("openpyxl", "openpyxl"),
        "beautifulsoup4": _dependency_status("beautifulsoup4", "bs4"),
        "pdf2image": _dependency_status("pdf2image", "pdf2image"),
        "pillow": _dependency_status("pillow", "PIL"),
        "pytesseract": _dependency_status("pytesseract", "pytesseract"),
    }

    tesseract_cmd = _resolve_tesseract_cmd()
    tessdata_path = _resolve_tessdata_path(Path(tesseract_cmd) if tesseract_cmd else None)
    languages = _available_tessdata_languages(tessdata_path)

    poppler_path = _resolve_poppler_path()
    pdftoppm = None
    if poppler_path:
        executable = "pdftoppm.exe" if os.name == "nt" else "pdftoppm"
        candidate = Path(poppler_path) / executable
        if candidate.exists():
            pdftoppm = str(candidate)

    ocr_ready = all([
        dependencies["pdf2image"]["available"],
        dependencies["pillow"]["available"],
        dependencies["pytesseract"]["available"],
        bool(tesseract_cmd),
        bool(pdftoppm),
        bool(languages),
    ])

    return {
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
        "text_extensions": sorted(TEXT_EXTENSIONS),
        "image_extensions": sorted(IMAGE_EXTENSIONS),
        "dependencies": dependencies,
        "ocr": {
            "ready": ocr_ready,
            "backend": "tesseract",
            "tesseract": _executable_version(tesseract_cmd, "--version"),
            "poppler": _executable_version(pdftoppm, "-v"),
            "tessdata_path": tessdata_path,
            "language_count": len(languages),
            "common_languages": {
                "eng": "eng" in languages,
                "chi_sim": "chi_sim" in languages,
                "chi_tra": "chi_tra" in languages,
                "osd": "osd" in languages,
            },
        },
    }


def _ocr_pdf_pages(content: bytes, page_numbers: List[int], options: LoaderOptions) -> Dict[int, str]:
    if options.ocr_max_pages is not None:
        page_numbers = page_numbers[:options.ocr_max_pages]

    if options.ocr_static_pages is not None:
        return {
            page_number: options.ocr_static_pages.get(page_number, "").strip()
            for page_number in page_numbers
            if options.ocr_static_pages.get(page_number, "").strip()
        }

    if options.ocr_backend.lower() != "tesseract":
        raise ValueError(f"Unsupported OCR backend: {options.ocr_backend}")

    try:
        from pdf2image import convert_from_bytes
        import pytesseract
    except ImportError:
        raise ImportError(
            "PDF OCR requires pdf2image and pytesseract. Install them and ensure Poppler/Tesseract are available."
        )

    _configure_tesseract(pytesseract)
    poppler_path = _resolve_poppler_path()
    ocr_results: Dict[int, str] = {}
    for page_number in page_numbers:
        convert_kwargs = {
            "dpi": options.ocr_dpi,
            "first_page": page_number,
            "last_page": page_number,
        }
        if poppler_path:
            convert_kwargs["poppler_path"] = poppler_path
        images = convert_from_bytes(content, **convert_kwargs)
        if not images:
            continue
        text = pytesseract.image_to_string(images[0], lang=options.ocr_language).strip()
        if text:
            ocr_results[page_number] = text
    return ocr_results


def _ocr_image(content: bytes, options: LoaderOptions) -> Dict[int, str]:
    if options.ocr_static_pages is not None:
        text = options.ocr_static_pages.get(1, "").strip()
        return {1: text} if text else {}

    if options.ocr_backend.lower() != "tesseract":
        raise ValueError(f"Unsupported OCR backend: {options.ocr_backend}")

    try:
        from PIL import Image
        import pytesseract
    except ImportError:
        raise ImportError("Image OCR requires pillow and pytesseract: pip install pillow pytesseract")

    _configure_tesseract(pytesseract)
    image = Image.open(io.BytesIO(content))
    text = pytesseract.image_to_string(image, lang=options.ocr_language).strip()
    return {1: text} if text else {}


def _infer_section_title(text: str) -> Optional[str]:
    """Infer a stable section title from the first heading-like line on a PDF page."""
    for line in text.splitlines():
        candidate = re.sub(r"\s+", " ", line).strip()
        if not candidate:
            continue
        if len(candidate) > 120:
            continue
        if len(candidate.split()) > 14:
            continue
        if re.fullmatch(r"[\W\d_]+", candidate):
            continue
        if candidate.endswith((".", ",", ";", ":", "，", "。", "；")) and len(candidate) > 30:
            continue
        return candidate
    return None


def _docx_heading_level(paragraph: Any) -> Optional[int]:
    style_name = ""
    if paragraph.style is not None:
        style_name = paragraph.style.name or ""

    match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
    if not match:
        return None

    level = int(match.group(1))
    if level < 1 or level > 9:
        return None
    return level


def _excel_rows_with_merged_values(sheet: Any) -> List[List[str]]:
    merged_values: Dict[Tuple[int, int], Any] = {}
    for merged_range in sheet.merged_cells.ranges:
        anchor = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row_number in range(merged_range.min_row, merged_range.max_row + 1):
            for column_number in range(merged_range.min_col, merged_range.max_col + 1):
                merged_values[(row_number, column_number)] = anchor

    rows: List[List[str]] = []
    for row_number in range(1, sheet.max_row + 1):
        values = []
        for column_number in range(1, sheet.max_column + 1):
            value = sheet.cell(row_number, column_number).value
            if value is None:
                value = merged_values.get((row_number, column_number))
            values.append(_cell_to_text(value))
        rows.append(values)

    return rows


def _excel_header_start_row(rows: List[List[str]], options: LoaderOptions) -> Optional[int]:
    if options.excel_header_row is not None:
        if 1 <= options.excel_header_row <= len(rows):
            return options.excel_header_row
        return None

    for index, row in enumerate(rows, start=1):
        if any(row):
            return index
    return None


def _combine_excel_headers(header_rows: List[List[str]]) -> List[str]:
    if not header_rows:
        return []

    width = max(len(row) for row in header_rows)
    headers = []
    for column_index in range(width):
        parts = []
        for row in header_rows:
            value = row[column_index] if column_index < len(row) else ""
            if value and value not in parts:
                parts.append(value)
        headers.append(" / ".join(parts) if parts else _excel_column_name(column_index + 1))
    return headers


def _excel_row_to_text(
    sheet_name: str,
    row_number: int,
    headers: List[str],
    row_values: List[str]
) -> str:
    lines = [f"Sheet: {sheet_name}", f"Row: {row_number}"]
    for column_index, value in enumerate(row_values):
        if not value:
            continue
        header = headers[column_index] if column_index < len(headers) else _excel_column_name(column_index + 1)
        lines.append(f"{header}: {value}")
    return "\n".join(lines)


def _excel_merged_ranges(sheet: Any) -> List[str]:
    return [str(merged_range) for merged_range in sheet.merged_cells.ranges]


def _excel_column_name(column_number: int) -> str:
    name = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        name = chr(65 + remainder) + name
    return f"Column {name}"


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def _docx_table_to_markdown(table: Any, table_index: int) -> str:
    rows = []
    for row in table.rows:
        values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(values):
            rows.append(values)

    if not rows:
        return ""

    header = rows[0]
    separators = ["---" for _ in header]
    body = rows[1:]
    lines = [f"Table {table_index}:", _markdown_row(header), _markdown_row(separators)]
    lines.extend(_markdown_row(row) for row in body)
    return "\n".join(lines)


def _markdown_row(values: List[str]) -> str:
    return "| " + " | ".join(values) + " |"
